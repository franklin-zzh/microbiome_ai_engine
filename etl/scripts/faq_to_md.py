"""FAQ 问答对 → qa_model 入库 md（Task 2.2 Q-to-Q）。

输入（三选一）：
- CSV：utf-8-sig（Excel 直接另存可得），表头「问题,答案」；
- XLSX：Sheet 1 表头「问题/答案」（需 openpyxl）；
- MD：连续「问题：… / 答案：…」段落。

输出：qa_model 格式 md（每对 `问题：…\\n\\n答案：…`，与 Dify QA Processor 的识别格式一致），
      再由 import_etl_md.py --doc-form qa_model 经 Knowledge Pipeline 入库 CS_QA（cs_qa）。

用法（项目根目录）：
    etl\\.venv\\Scripts\\python.exe scripts\\faq_to_md.py -i faq.csv -o out\\faq_qa.md --title "客服高频问答"
    etl\\.venv\\Scripts\\python.exe scripts\\faq_to_md.py -i faq.xlsx -o out\\faq_qa.md
"""
import argparse
import csv
import sys
from pathlib import Path


def load_csv(path: Path) -> list[tuple[str, str]]:
    pairs: list[tuple[str, str]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        header = next(reader, None)
        start = 1 if header and header[0].strip().lower() in {"问题", "question", "q"} else 0
        for row in reader:
            if len(row) >= 2 and row[0].strip() and row[1].strip():
                pairs.append((row[0].strip(), row[1].strip()))
    return pairs


def load_xlsx(path: Path) -> list[tuple[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    ws = wb.active
    pairs: list[tuple[str, str]] = []
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    start = 1 if header and str(header[0] or "").strip() in {"问题", "question", "q"} else 0
    for row in rows:
        if row and str(row[0] or "").strip() and len(row) >= 2 and str(row[1] or "").strip():
            pairs.append((str(row[0]).strip(), str(row[1]).strip()))
    wb.close()
    return pairs


def load_md(path: Path) -> list[tuple[str, str]]:
    """解析「问题：… / 答案：…」段落对。"""
    text = path.read_text(encoding="utf-8-sig")  # utf-8-sig 兼容 Windows 记事本/Excel 导出的 BOM
    pairs: list[tuple[str, str]] = []
    question = answer = None
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("问题："):
            if question and answer:
                pairs.append((question, answer))
            question = stripped[len("问题：") :].strip()
            answer = None
        elif stripped.startswith("答案：") and question is not None:
            answer = stripped[len("答案：") :].strip()
        elif question is not None and answer is not None:
            answer += " " + stripped
    if question and answer:
        pairs.append((question, answer))
    return pairs


def to_qa_md(pairs: list[tuple[str, str]], title: str = "") -> str:
    blocks = [f"# {title}\n" if title else ""]
    for question, answer in pairs:
        blocks.append(f"问题：{question}\n\n答案：{answer}\n")
    return "\n".join(b for b in blocks if b)


def main() -> int:
    parser = argparse.ArgumentParser(description="FAQ 问答对 → qa_model 入库 md（Q-to-Q）")
    parser.add_argument("-i", "--input", required=True, help="输入文件：CSV / XLSX / MD（问题：/答案：对）")
    parser.add_argument("-o", "--output", required=True, help="输出 qa_model md 路径")
    parser.add_argument("--title", default="客服 FAQ 问答对", help="输出 md 标题（默认：客服 FAQ 问答对）")
    args = parser.parse_args()

    src = Path(args.input)
    if not src.is_file():
        print(f"[ERROR] 输入文件不存在: {src}")
        return 1
    ext = src.suffix.lower()
    try:
        if ext == ".csv":
            pairs = load_csv(src)
        elif ext == ".xlsx":
            pairs = load_xlsx(src)
        elif ext in {".md", ".txt"}:
            pairs = load_md(src)
        else:
            print(f"[ERROR] 不支持的输入格式: {ext}（支持 .csv/.xlsx/.md/.txt）")
            return 1
    except Exception as exc:
        print(f"[ERROR] 解析失败: {exc}")
        return 1

    if not pairs:
        print("[ERROR] 未解析到任何 问题/答案 对（表头应为「问题,答案」或「问题：/答案：」）")
        return 1

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(to_qa_md(pairs, args.title), encoding="utf-8")
    print(f"[OK] {len(pairs)} 对问答已写入 {out}")
    print("入库：import_etl_md.py --md-dir <输出目录> --doc-form qa_model")
    return 0


if __name__ == "__main__":
    sys.exit(main())
