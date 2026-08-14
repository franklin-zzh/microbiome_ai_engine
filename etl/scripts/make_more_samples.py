"""生成 PDF（原生+扫描件）与 DOCX 冒烟样例（需已安装 etl/requirements.txt）。"""
import sys
from pathlib import Path


def main(out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)

    import pymupdf  # fitz 旧别名已废弃

    # 原生 PDF（含文本层）：insert_text 需指定中文字体，否则字符不嵌入（提取出来是乱码）
    fontfile = next(
        (p for p in ("C:/Windows/Fonts/simhei.ttf", "C:/Windows/Fonts/simsun.ttc", "C:/Windows/Fonts/msyh.ttc") if Path(p).exists()),
        None,
    )
    doc = pymupdf.open()
    page = doc.new_page()
    if fontfile:
        page.insert_text((72, 72), "Native PDF: 肠道菌群检测说明", fontsize=14, fontfile=fontfile, fontname="cnfont")
        page.insert_text((72, 100), "本 PDF 含文本层，可直接提取。采样后 7 个工作日出具报告。", fontsize=11, fontfile=fontfile, fontname="cnfont")
    else:
        page.insert_text((72, 72), "Native PDF with text layer.", fontsize=14)
    doc.save(str(out_dir / "native.pdf"))
    doc.close()

    # 扫描件 PDF（无文本层：整页是一张图片）
    img = out_dir / "chart_text.png"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, filename=str(img))
    doc.save(str(out_dir / "scanned.pdf"))
    doc.close()

    # DOCX
    from docx import Document

    d = Document()
    d.add_heading("DOCX 测试文档", level=1)
    d.add_paragraph("这是一个 Word 文档，用于验证 markitdown 转换。")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "采样"
    table.cell(1, 1).text = "粪便采样盒"
    d.save(str(out_dir / "test.docx"))

    print(f"saved {out_dir}")


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else Path("out"))
