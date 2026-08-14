"""PPTX → Block：python-pptx 逐页结构提取（不走 PDF 中转）。

- 标题（slide.shapes.title）→ heading3；正文文本框 → 段落（按 paragraph.level 提升为小标题）；
- 表格（shape.has_table）→ table 块；
- 图片（shape_type=PICTURE）→ 抠出保存到 workdir，交由 OCR 阶段识别图片内文字；
- GROUP 组形状递归处理；
- 演讲者备注（notes）→ 引用块（备注常含讲解要点，不应丢弃）。
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from etl.blocks import Block


def _iter_text_paragraphs(text_frame):
    """按段落产出 (level, text)，过滤空段落。"""
    for para in text_frame.paragraphs:
        text = "".join(run.text for run in para.runs).strip()
        if not text:
            continue
        yield para.level, text


def extract(path: Path, workdir: Path) -> list[Block]:
    workdir.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(path))
    blocks: list[Block] = []
    pic_seq = 0

    def handle_shape(shape, page_index: int) -> None:
        nonlocal pic_seq
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            pic_seq += 1
            image = shape.image
            ext = f".{image.ext}" if image.ext else ".png"
            img_path = workdir / f"page{page_index:03d}-pic{pic_seq:02d}{ext}"
            try:
                img_path.write_bytes(image.blob)
            except OSError:
                return
            blocks.append(
                Block(
                    kind="image",
                    image_path=img_path,
                    caption=f"第{page_index}页图片{pic_seq}",
                    meta={"page": page_index},
                )
            )
        elif shape.has_table:
            table = shape.table
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if any(any(c for c in row) for row in rows):
                blocks.append(Block(kind="table", rows=rows, meta={"page": page_index}))
        elif shape.has_text_frame:
            for level, text in _iter_text_paragraphs(shape.text_frame):
                if level >= 2:
                    blocks.append(Block(kind="heading", level=min(5, 3 + level - 2), text=text))
                else:
                    blocks.append(Block(kind="paragraph", text=text, meta={"page": page_index}))
        elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub in shape.shapes:
                handle_shape(sub, page_index)

    for page_index, slide in enumerate(prs.slides, start=1):
        blocks.append(Block(kind="heading", level=2, text=f"第 {page_index} 页"))
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title_text = slide.shapes.title.text_frame.text.strip()
            if title_text:
                blocks.append(Block(kind="heading", level=3, text=title_text))
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            handle_shape(shape, page_index)
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                blocks.append(Block(kind="quote", text=notes_text, meta={"page": page_index}))

    return blocks
